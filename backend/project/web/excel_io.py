"""
Excel export/import for products and sales (openpyxl .xlsx).
"""
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib.auth.models import User
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from api.models import Category, Product, Sale, SaleItem

PRODUCT_HEADERS = [
    'name',
    'barcode',
    'category',
    'price',
    'stock_quantity',
    'low_stock_threshold',
    'description',
]

SALE_ITEM_HEADERS = [
    'sale_reference',
    'sale_date',
    'cashier_username',
    'product_barcode',
    'quantity',
    'unit_price',
    'line_subtotal',
    'sale_discount',
    'sale_grand_total',
]


def _style_header_row(ws, headers):
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='2563EB')
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill


def workbook_to_bytes(wb):
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_products_workbook(products):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    _style_header_row(ws, PRODUCT_HEADERS)

    for row_idx, product in enumerate(products, start=2):
        ws.cell(row=row_idx, column=1, value=product.name)
        ws.cell(row=row_idx, column=2, value=product.barcode)
        ws.cell(row=row_idx, column=3, value=product.category.name if product.category else '')
        ws.cell(row=row_idx, column=4, value=float(product.price))
        ws.cell(row=row_idx, column=5, value=product.stock_quantity)
        ws.cell(row=row_idx, column=6, value=product.low_stock_threshold)
        ws.cell(row=row_idx, column=7, value=product.description or '')

    for col in range(1, len(PRODUCT_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    return wb


def build_products_template_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    _style_header_row(ws, PRODUCT_HEADERS)
    ws.append([
        'Sample Product',
        '1234567890123',
        'General',
        1500.00,
        50,
        5,
        'Optional description',
    ])
    note = wb.create_sheet('Instructions')
    note['A1'] = 'Required columns: name, barcode, price'
    note['A2'] = 'Import updates existing products by matching barcode.'
    note['A3'] = 'category will be created if it does not exist.'
    return wb


def build_sales_workbook(sales_qs):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sale Line Items'
    _style_header_row(ws, SALE_ITEM_HEADERS)

    row_idx = 2
    for sale in sales_qs.prefetch_related('items__product', 'cashier'):
        cashier = sale.cashier.username if sale.cashier else ''
        sale_ref = str(sale.id)
        sale_date = timezone.localtime(sale.created_at).strftime('%Y-%m-%d %H:%M')
        for item in sale.items.all():
            ws.cell(row=row_idx, column=1, value=sale_ref)
            ws.cell(row=row_idx, column=2, value=sale_date)
            ws.cell(row=row_idx, column=3, value=cashier)
            ws.cell(row=row_idx, column=4, value=item.product.barcode)
            ws.cell(row=row_idx, column=5, value=item.quantity)
            ws.cell(row=row_idx, column=6, value=float(item.unit_price))
            ws.cell(row=row_idx, column=7, value=float(item.subtotal))
            ws.cell(row=row_idx, column=8, value=float(sale.discount))
            ws.cell(row=row_idx, column=9, value=float(sale.grand_total))
            row_idx += 1

    summary = wb.create_sheet('Sales Summary')
    summary_headers = ['sale_id', 'date', 'cashier', 'items_count', 'total_amount', 'discount', 'grand_total']
    _style_header_row(summary, summary_headers)
    s_row = 2
    for sale in sales_qs:
        summary.cell(row=s_row, column=1, value=str(sale.id))
        summary.cell(row=s_row, column=2, value=timezone.localtime(sale.created_at).strftime('%Y-%m-%d %H:%M'))
        summary.cell(row=s_row, column=3, value=sale.cashier.username if sale.cashier else '')
        summary.cell(row=s_row, column=4, value=sale.items.count())
        summary.cell(row=s_row, column=5, value=float(sale.total_amount))
        summary.cell(row=s_row, column=6, value=float(sale.discount))
        summary.cell(row=s_row, column=7, value=float(sale.grand_total))
        s_row += 1

    return wb


def build_sales_template_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sale Line Items'
    _style_header_row(ws, SALE_ITEM_HEADERS)
    ws.append([
        'SALE-001',
        '2026-05-26 10:00',
        'admin',
        '1234567890123',
        2,
        500.00,
        1000.00,
        0.00,
        1000.00,
    ])
    note = wb.create_sheet('Instructions')
    note['A1'] = 'Group rows with the same sale_reference to create one sale.'
    note['A2'] = 'product_barcode must exist in inventory.'
    note['A3'] = 'sale_grand_total and sale_discount apply once per sale (first row).'
    return wb


def _cell_str(value):
    if value is None:
        return ''
    return str(value).strip()


def _parse_decimal(value, default=Decimal('0')):
    if value is None or value == '':
        return default
    try:
        return Decimal(str(value).replace(',', ''))
    except (InvalidOperation, ValueError):
        raise ValueError(f'Invalid number: {value}')


def _parse_int(value, default=0):
    if value is None or value == '':
        return default
    return int(float(str(value)))


def _get_or_create_category(name):
    if not name:
        return None
    category, _ = Category.objects.get_or_create(name=name[:100])
    return category


def import_products_from_file(uploaded_file):
    wb = load_workbook(uploaded_file, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {'created': 0, 'updated': 0, 'errors': ['File is empty.']}

    headers = [_cell_str(h).lower() for h in rows[0]]
    col_map = {h: i for i, h in enumerate(headers) if h}

    required = {'name', 'barcode', 'price'}
    if not required.issubset(col_map.keys()):
        return {'created': 0, 'updated': 0, 'errors': [f'Missing columns. Required: {", ".join(required)}']}

    created = updated = 0
    errors = []

    for row_num, row in enumerate(rows[1:], start=2):
        if not row or all(v is None or _cell_str(v) == '' for v in row):
            continue

        def val(key):
            idx = col_map.get(key)
            if idx is None or idx >= len(row):
                return ''
            return _cell_str(row[idx])

        name = val('name')
        barcode = val('barcode')
        if not name or not barcode:
            errors.append(f'Row {row_num}: name and barcode are required.')
            continue

        try:
            price = _parse_decimal(val('price'))
            stock = _parse_int(val('stock_quantity'), 0)
            threshold = _parse_int(val('low_stock_threshold'), 5)
            category = _get_or_create_category(val('category'))
            description = val('description') or None

            product, was_created = Product.objects.update_or_create(
                barcode=barcode,
                defaults={
                    'name': name,
                    'price': price,
                    'stock_quantity': stock,
                    'low_stock_threshold': threshold,
                    'category': category,
                    'description': description,
                },
            )
            if was_created:
                created += 1
            else:
                updated += 1
        except Exception as exc:
            errors.append(f'Row {row_num}: {exc}')

    return {'created': created, 'updated': updated, 'errors': errors}


def import_sales_from_file(uploaded_file, default_cashier=None):
    wb = load_workbook(uploaded_file, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {'sales_created': 0, 'errors': ['File is empty.']}

    headers = [_cell_str(h).lower() for h in rows[0]]
    col_map = {h: i for i, h in enumerate(headers) if h}

    if 'sale_reference' not in col_map or 'product_barcode' not in col_map:
        return {'sales_created': 0, 'errors': ['Missing sale_reference or product_barcode column.']}

    grouped = defaultdict(list)

    for row_num, row in enumerate(rows[1:], start=2):
        if not row or all(v is None or _cell_str(v) == '' for v in row):
            continue

        def val(key):
            idx = col_map.get(key)
            if idx is None or idx >= len(row):
                return ''
            return _cell_str(row[idx])

        ref = val('sale_reference') or f'ROW-{row_num}'
        grouped[ref].append((row_num, val))

    sales_created = 0
    errors = []

    with transaction.atomic():
        for ref, line_rows in grouped.items():
            try:
                first_row_num, first = line_rows[0]
                cashier_username = first('cashier_username')
                cashier = default_cashier
                if cashier_username:
                    cashier = User.objects.filter(username=cashier_username).first()
                    if not cashier:
                        errors.append(f'Sale {ref}: cashier "{cashier_username}" not found.')
                        continue

                discount = _parse_decimal(first('sale_discount'), Decimal('0'))
                grand_total = _parse_decimal(first('sale_grand_total'), Decimal('0'))

                items_data = []
                computed_total = Decimal('0')

                for row_num, row_vals in line_rows:
                    barcode = row_vals('product_barcode')
                    product = Product.objects.filter(barcode=barcode).first()
                    if not product:
                        errors.append(f'Row {row_num}: product barcode "{barcode}" not found.')
                        continue

                    qty = _parse_int(row_vals('quantity'), 1)
                    if qty <= 0:
                        errors.append(f'Row {row_num}: quantity must be positive.')
                        continue

                    unit_price = _parse_decimal(row_vals('unit_price'), product.price)
                    line_sub = _parse_decimal(row_vals('line_subtotal'), unit_price * qty)
                    computed_total += line_sub
                    items_data.append({
                        'product': product,
                        'quantity': qty,
                        'unit_price': unit_price,
                        'subtotal': line_sub,
                    })

                if not items_data:
                    errors.append(f'Sale {ref}: no valid line items.')
                    continue

                if grand_total == 0:
                    grand_total = computed_total
                if grand_total == 0 and computed_total > 0:
                    grand_total = computed_total
                total_amount = computed_total

                sale = Sale.objects.create(
                    cashier=cashier,
                    total_amount=total_amount,
                    discount=discount,
                    grand_total=grand_total,
                    client_id=f'import-{ref}',
                )
                for item in items_data:
                    SaleItem.objects.create(
                        sale=sale,
                        product=item['product'],
                        quantity=item['quantity'],
                        unit_price=item['unit_price'],
                        subtotal=item['subtotal'],
                    )
                    item['product'].stock_quantity -= item['quantity']
                    item['product'].save(update_fields=['stock_quantity'])

                sales_created += 1
            except Exception as exc:
                errors.append(f'Sale {ref}: {exc}')

    return {'sales_created': sales_created, 'errors': errors}
